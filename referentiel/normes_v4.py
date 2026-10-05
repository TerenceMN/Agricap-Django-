"""
Normes institutionnelles du cycle de crédit → `InstitutionConfig`.

FICHIER SOURCE : `AGRICAP_FIN_Simulateur_Credit_Cycle_Production_v4.xlsx`, le
« fichier 2 » du moteur d'analyse — celui qui dit **à quoi comparer**, quand les
simulateurs par filière (`AGRICAP_FIN_SIM_NN_…`) disent **comment lire**.

POURQUOI CE MODULE EXISTE
-------------------------
Le classeur v4 était ingéré dans `dataio` puis ignoré : la commande
`ingest_simulateurs` l'exclut explicitement de son motif, au motif qu'il
« n'alimente aucun référentiel ». C'était vrai de `ReferentielFiliere` — le v4 ne
porte qu'un cycle maïs générique, il ne peut pas servir de référence de coûts à
quatorze filières. Ce n'était pas vrai de la GOUVERNANCE : ses onglets 1, 10 et
12 portent exactement les paramètres que le principe 8 exige de sortir du code.

Ces valeurs vivaient en dur dans `credits.analyse` :

    POIDS_DEFAUT              = 25 / 20 / 10 / 30 / 15   ← onglet 12
    REGLES_DECISION_DEFAUT    = DSCR 1,2 puis 1,0        ← onglet 1
    (frais de dossier, commissions)                      ← onglet 10

Le comité ne pouvait donc pas recalibrer sa propre grille sans un déploiement.
Il le peut désormais en re-téléversant son classeur — et l'écart entre deux
versions est lisible, puisque chaque config conserve dans `raw` les lignes
sources qui l'ont produite.

CE QUE CE MODULE NE FAIT PAS
----------------------------
Il ne touche ni `ReferentielVersion` ni `ReferenceRange`. Le v4 n'est pas le
référentiel technico-économique v3 : créer une version à son nom désactiverait la
version v3 active, et avec elle toutes les plages de vraisemblance rattachées.
La configuration produite reste donc rattachée à la version v3 courante, et la
provenance du v4 est tracée dans `raw`.

Il n'écrit AUCUNE valeur absente du classeur. La configuration produite est une
COPIE de la configuration active, sur laquelle seules les valeurs effectivement
lues sont posées : les paramètres de caution solidaire, le plafond de délégation
et la phase de déploiement — qui ne figurent dans aucun onglet du v4 — survivent
à l'ingestion au lieu de retomber sur les défauts du modèle.
"""
from __future__ import annotations

import logging
import unicodedata
from dataclasses import dataclass, field
from decimal import Decimal

import openpyxl

from .ingest import activer_config
from .models import InstitutionConfig
from .range_parser import to_decimal

logger = logging.getLogger(__name__)

#: Nom canonique du classeur — la commande d'ingestion le cherche sous ce nom.
NOM_CLASSEUR = "AGRICAP_FIN_Simulateur_Credit_Cycle_Production_v4.xlsx"

ONGLET_PARAMETRES = "1_Accueil_Parametres"
ONGLET_CAPACITE = "10_Capacite_Remboursement"
ONGLET_SCORING = "12_Analyse_Credit"

#: Base 100 des pondérations. Le classeur les exprime en FRACTIONS (0,25) ;
#: `InstitutionConfig` les stocke en points (25,00), parce que
#: `credits.analyse.poids_effectifs()` vérifie que leur somme vaut exactement 100
#: avant de les retenir. Écrire 0,25 en base ferait échouer cet invariant à
#: chaque analyse et retomber sur les poids de secours — c'est-à-dire annuler
#: silencieusement le recalibrage que le comité vient de voter.
CENT = Decimal(100)


def _norm(text) -> str:
    """Minuscule, sans accent, espaces compactés — même normalisation que
    `referentiel.ingest`, pour que les deux lecteurs de la même table de
    paramètres reconnaissent les mêmes libellés."""
    if text is None:
        return ""
    s = unicodedata.normalize("NFD", str(text))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return " ".join(s.lower().split())


#: Libellé (fragment normalisé) → attribut `InstitutionConfig`, onglet 1.
#:
#: Les fragments évitent volontairement les apostrophes : le classeur mélange
#: l'apostrophe droite (') et la typographique (’) selon les cellules, et un
#: fragment « taux d'interet » ne reconnaîtrait qu'une moitié des versions du
#: fichier. « interet annuel » les reconnaît toutes.
NORMES_PARAMETRES: dict[str, str] = {
    "seuil dscr (avis": "seuil_dscr",
    "seuil dscr stresse": "seuil_dscr_stresse",
    "seuil de couverture des garanties": "couverture_min",
    "score global minimal": "score_global_min",
    "interet annuel": "taux_interet_annuel",
}

#: Onglet 10 — les deux taux qui composent le coût du crédit annoncé au client.
#: « frais de dossier » et non « frais » : la feuille porte aussi « Frais totaux »,
#: qui est un MONTANT calculé et non un taux.
NORMES_CAPACITE: dict[str, str] = {
    "frais de dossier": "frais_dossier",
    "commissions": "commissions",
}

#: Onglet 12 — les cinq pondérations, dans l'ordre de la grille.
#: « garanties & domiciliation » et non « garanties » : la même feuille porte
#: « Couverture par garanties (info) », qui est un ratio de dossier et non un poids.
NORMES_PONDERATIONS: dict[str, str] = {
    "fiabilite technique": "poids_technique",
    "capacite financiere": "poids_financier",
    "resilience au stress": "poids_stress",
    "historique comportemental": "poids_comportemental",
    "garanties & domiciliation": "poids_garanties",
}

#: Les cinq attributs de pondération, traités comme un BLOC indivisible.
ATTRS_POIDS = tuple(NORMES_PONDERATIONS.values())

#: Bornes de plausibilité — un garde-fou de saisie, pas un seuil métier
#: (principe 8). Elles n'expriment aucune politique de crédit : elles disent
#: seulement qu'une cellule hors de ces bornes est une erreur de saisie ou une
#: colonne mal lue, jamais une décision du comité. Un taux d'intérêt annuel à 18
#: (points) au lieu de 0,18 (fraction) est le cas réel qu'elles attrapent.
BORNES: dict[str, tuple[Decimal, Decimal]] = {
    "seuil_dscr": (Decimal("0"), Decimal("10")),
    "seuil_dscr_stresse": (Decimal("0"), Decimal("10")),
    "couverture_min": (Decimal("0"), Decimal("10")),
    "score_global_min": (Decimal("0"), Decimal("100")),
    "taux_interet_annuel": (Decimal("0"), Decimal("1")),
    "frais_dossier": (Decimal("0"), Decimal("1")),
    "commissions": (Decimal("0"), Decimal("1")),
}

#: Quantize par famille — reprend à l'identique les précisions documentées sur
#: `InstitutionConfig` (ratios 0,001 ; poids et scores 0,01 ; fractions 0,0001).
QUANTIZE: dict[str, Decimal] = {
    "seuil_dscr": Decimal("0.001"),
    "seuil_dscr_stresse": Decimal("0.001"),
    "couverture_min": Decimal("0.001"),
    "score_global_min": Decimal("0.01"),
    "taux_interet_annuel": Decimal("0.0001"),
    "frais_dossier": Decimal("0.0001"),
    "commissions": Decimal("0.0001"),
    **{attr: Decimal("0.01") for attr in ATTRS_POIDS},
}


class NormesIntrouvables(Exception):
    """Le classeur ne porte aucun des onglets de normes attendus."""


@dataclass
class RapportNormes:
    """Ce qui a été lu, ce qui ne l'a pas été, et pourquoi."""

    valeurs: dict[str, Decimal] = field(default_factory=dict)
    onglets_lus: list[str] = field(default_factory=list)
    onglets_absents: list[str] = field(default_factory=list)
    avertissements: list[str] = field(default_factory=list)
    lignes_source: dict[str, list[list[str]]] = field(default_factory=dict)
    config_id: int | None = None

    @property
    def poids_appliques(self) -> bool:
        return all(attr in self.valeurs for attr in ATTRS_POIDS)


# ── Lecture ──────────────────────────────────────────────────────────────────

def _paires(ws) -> list[tuple[str, object, list[str]]]:
    """(libellé normalisé, valeur brute colonne B, ligne source) de chaque ligne."""
    sorties = []
    for row in ws.iter_rows(values_only=True):
        if not row:
            continue
        brutes = [("" if c is None else str(c)) for c in row]
        if not any(c.strip() for c in brutes):
            continue
        valeur = row[1] if len(row) > 1 else None
        sorties.append((_norm(row[0]), valeur, [c for c in brutes if c.strip()]))
    return sorties


def _lire_onglet(ws, mapping: dict[str, str], rapport: RapportNormes,
                 *, facteur: Decimal = Decimal(1)) -> dict[str, Decimal]:
    """Applique `mapping` aux lignes de `ws`. Une cellule illisible est SIGNALÉE.

    Le premier fragment qui correspond gagne, et la ligne est consommée : les
    mappings sont écrits pour être mutuellement exclusifs sur un même onglet.
    """
    trouvees: dict[str, Decimal] = {}
    lignes: list[list[str]] = []

    for label, valeur, source in _paires(ws):
        lignes.append(source)
        for fragment, attr in mapping.items():
            if fragment not in label:
                continue
            if attr in trouvees:
                rapport.avertissements.append(
                    f"« {source[0]} » : le paramètre « {attr} » avait déjà été lu "
                    f"plus haut dans l'onglet — la première occurrence est "
                    f"conservée, celle-ci est ignorée.")
                break

            brut = to_decimal(valeur)
            if brut is None:
                rapport.avertissements.append(
                    f"« {source[0]} » : valeur illisible ({valeur!r}) — le "
                    f"paramètre « {attr} » garde sa valeur actuelle.")
                break

            montant = (brut * facteur).quantize(QUANTIZE.get(attr, Decimal("0.001")))
            borne = BORNES.get(attr)
            if borne and not (borne[0] <= montant <= borne[1]):
                rapport.avertissements.append(
                    f"« {source[0]} » : valeur {montant} hors des bornes de "
                    f"plausibilité [{borne[0]} ; {borne[1]}] — paramètre "
                    f"« {attr} » NON repris. Vérifiez l'unité de la cellule "
                    f"(fraction attendue, pas des points de pourcentage).")
                break

            trouvees[attr] = montant
            break

    rapport.lignes_source[ws.title] = lignes
    return trouvees


def _valider_poids(poids: dict[str, Decimal], rapport: RapportNormes) -> dict[str, Decimal]:
    """Les cinq poids passent ou ne passent pas ENSEMBLE.

    Une pondération partielle est le pire des cas : trois poids repris du v4 et
    deux hérités de la config précédente donneraient une somme quelconque, et
    `poids_effectifs()` retomberait sur les poids de secours — l'analyste
    croirait scorer avec la grille du comité et scorerait avec celle du code.
    On préfère ne rien reprendre et le dire.
    """
    manquants = [a for a in ATTRS_POIDS if a not in poids]
    if manquants:
        rapport.avertissements.append(
            f"Pondérations non reprises : {len(manquants)} critère(s) sur 5 "
            f"introuvable(s) dans l'onglet « {ONGLET_SCORING} » "
            f"({', '.join(manquants)}). La grille de pondération actuelle est "
            f"conservée en entier — une grille à moitié recalibrée ne somme plus "
            f"à 100 et serait ignorée par le moteur.")
        return {}

    total = sum(poids.values(), Decimal(0))
    if total != CENT:
        rapport.avertissements.append(
            f"Pondérations non reprises : leur somme vaut {total} et non 100. "
            f"Le moteur exige 100 pour que le score global soit lisible "
            f"(« 68/100 » sur une base de 100). Corrigez l'onglet "
            f"« {ONGLET_SCORING} » puis re-téléversez.")
        return {}

    return poids


def lire_normes(path_or_file) -> RapportNormes:
    """Lit les normes du classeur v4. **Aucune écriture en base.**

    C'est le mode que la commande expose en `--check` : l'admin voit ce qui
    serait repris AVANT de remplacer la configuration de l'institution.
    """
    wb = openpyxl.load_workbook(path_or_file, data_only=True, read_only=True)
    rapport = RapportNormes()
    try:
        valeurs: dict[str, Decimal] = {}

        for onglet, mapping, facteur in (
            (ONGLET_PARAMETRES, NORMES_PARAMETRES, Decimal(1)),
            (ONGLET_CAPACITE, NORMES_CAPACITE, Decimal(1)),
        ):
            if onglet not in wb.sheetnames:
                rapport.onglets_absents.append(onglet)
                continue
            rapport.onglets_lus.append(onglet)
            valeurs.update(_lire_onglet(wb[onglet], mapping, rapport, facteur=facteur))

        if ONGLET_SCORING in wb.sheetnames:
            rapport.onglets_lus.append(ONGLET_SCORING)
            poids = _lire_onglet(wb[ONGLET_SCORING], NORMES_PONDERATIONS, rapport,
                                 facteur=CENT)
            valeurs.update(_valider_poids(poids, rapport))
        else:
            rapport.onglets_absents.append(ONGLET_SCORING)

        if not rapport.onglets_lus:
            raise NormesIntrouvables(
                f"Aucun onglet de normes dans ce classeur : ni "
                f"« {ONGLET_PARAMETRES} », ni « {ONGLET_CAPACITE} », ni "
                f"« {ONGLET_SCORING} ». Ce n'est pas le simulateur de cycle de "
                f"crédit ({NOM_CLASSEUR}).")

        rapport.valeurs = valeurs
        return rapport
    finally:
        wb.close()


# ── Écriture ─────────────────────────────────────────────────────────────────

def ingerer_normes(path_or_file, *, source_filename: str = "") -> RapportNormes:
    """Lit le classeur v4 et en fait la configuration ACTIVE de l'institution.

    La nouvelle configuration est une COPIE de l'active, surchargée des seules
    valeurs lues : ce que le v4 ne porte pas (caution solidaire, plafond de
    délégation, décotes, phase de déploiement) est conservé tel quel. Repartir
    d'un `InstitutionConfig()` neuf, comme le fait l'import du référentiel v3,
    remettrait ces paramètres aux défauts du modèle sans que personne ne l'ait
    demandé — un plafond de délégation ramené à 25 000 USD par un upload de
    normes de scoring est exactement le genre d'effet de bord qu'un auditeur
    reproche deux ans plus tard.
    """
    rapport = lire_normes(path_or_file)

    courante = InstitutionConfig.active()
    nouvelle = InstitutionConfig.objects.get(pk=courante.pk) if courante.pk else InstitutionConfig()
    nouvelle.pk = None
    nouvelle.id = None
    # La version reste celle du référentiel v3 en cours : le v4 ne fournit ni
    # plages ni chaînes, il n'a pas à ouvrir une version de référentiel.
    nouvelle.version = courante.version if courante.pk else None

    for attr, valeur in rapport.valeurs.items():
        setattr(nouvelle, attr, valeur)

    nouvelle.raw = {
        **(courante.raw or {}),
        "normes_v4": {
            "fichier": source_filename or NOM_CLASSEUR,
            "ongletsLus": rapport.onglets_lus,
            "ongletsAbsents": rapport.onglets_absents,
            "parametresRepris": {a: str(v) for a, v in sorted(rapport.valeurs.items())},
            "avertissements": rapport.avertissements,
            "lignesSource": rapport.lignes_source,
        },
    }

    activer_config(nouvelle)
    rapport.config_id = nouvelle.pk

    for message in rapport.avertissements:
        logger.warning("Normes v4 (%s) : %s", source_filename or NOM_CLASSEUR, message)
    logger.info(
        "Normes v4 (%s) : configuration #%s activée, %s paramètre(s) repris (%s).",
        source_filename or NOM_CLASSEUR, nouvelle.pk, len(rapport.valeurs),
        ", ".join(sorted(rapport.valeurs)) or "aucun")

    return rapport
